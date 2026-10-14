"""AI Resume ATS Checker - Streamlit + Google Gemini Flash.

Upload a resume (PDF / DOCX / TXT), optionally paste a job description,
and get an ATS score with concrete improvement suggestions.
"""

import io
import json
import os
import re
import time

import streamlit as st

DEFAULT_MODEL = "gemini-flash-latest"  # alias that always points to the newest Flash model
FALLBACK_MODELS = ["gemini-flash-lite-latest"]  # tried if the main model stays overloaded
RETRIES_PER_MODEL = 3
BACKOFF_SECONDS = 2  # waits 2s, 4s between retries
TRANSIENT_MARKERS = ("UNAVAILABLE", "RESOURCE_EXHAUSTED", "DEADLINE_EXCEEDED", "INTERNAL", "503", "429", "504")
MAX_RESUME_CHARS = 15000
MAX_JD_CHARS = 6000

SECTION_PATTERNS = {
    "Summary / Objective": r"\b(summary|objective|profile|about me)\b",
    "Experience": r"\b(experience|employment|work history|internship)\b",
    "Education": r"\b(education|academic|qualification)\b",
    "Skills": r"\b(skills|technologies|technical skills|competencies)\b",
    "Projects": r"\b(projects|portfolio)\b",
}


# --------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------
def extract_text(file_bytes: bytes, filename: str) -> str:
    """Return plain text from a PDF, DOCX or TXT file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(file_bytes))
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()
    if name.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # many resumes keep content in tables
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# --------------------------------------------------------------------------
# Rule-based checks (fast, deterministic, no API needed)
# --------------------------------------------------------------------------
def local_checks(text: str) -> dict:
    """Simple ATS-friendliness checks. Returns a score (0-100) and details."""
    lowered = text.lower()
    words = re.findall(r"\b\w+\b", text)
    word_count = len(words)

    checks = []  # (label, passed, points)

    has_email = bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text))
    checks.append(("Email address found", has_email, 15))

    has_phone = bool(re.search(r"(\+?\d[\d\s().-]{8,}\d)", text))
    checks.append(("Phone number found", has_phone, 10))

    has_link = "linkedin.com" in lowered or "github.com" in lowered
    checks.append(("LinkedIn / GitHub link found", has_link, 5))

    found_sections = [
        name for name, pat in SECTION_PATTERNS.items() if re.search(pat, lowered)
    ]
    for core in ("Experience", "Education", "Skills"):
        checks.append((f"'{core}' section detected", core in found_sections, 12))

    length_ok = 150 <= word_count <= 1200
    checks.append((f"Reasonable length ({word_count} words)", length_ok, 14))

    # Bullets: normal bullet characters at line start, or a lone symbol on its own
    # line (some PDF fonts extract bullets that way).
    bullet_count = len(re.findall(r"^\s*[-•*▪●◦]", text, flags=re.MULTILINE))
    bullet_count += len(re.findall(r"^\s*[^\w\s]\s*$", text, flags=re.MULTILINE))
    checks.append(("Uses bullet points", bullet_count >= 5, 8))

    has_numbers = len(re.findall(r"\b\d+(?:[.,]\d+)?\s*(?:%|\+|k|m|x)?", text)) >= 5
    checks.append(("Contains measurable numbers / metrics", has_numbers, 12))

    total = sum(points for _, _, points in checks)
    earned = sum(points for _, passed, points in checks if passed)
    score = round(100 * earned / total) if total else 0

    return {
        "score": score,
        "checks": [{"label": l, "passed": p} for l, p, _ in checks],
        "word_count": word_count,
        "sections_found": found_sections,
    }


# --------------------------------------------------------------------------
# Gemini
# --------------------------------------------------------------------------
def build_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n{job_description[:MAX_JD_CHARS]}"
        if job_description.strip()
        else "JOB DESCRIPTION: (not provided - evaluate for general ATS quality)"
    )
    return f"""You are an expert ATS (Applicant Tracking System) analyst and professional resume coach.
Evaluate the resume below{' against the job description' if job_description.strip() else ''}.
Be honest and specific. Never invent facts that are not in the resume.

Return ONLY a JSON object with exactly this structure:
{{
  "overall_score": <integer 0-100>,
  "category_scores": {{
    "keywords_match": <integer 0-100>,
    "formatting_structure": <integer 0-100>,
    "experience_impact": <integer 0-100>,
    "skills_relevance": <integer 0-100>,
    "readability_clarity": <integer 0-100>
  }},
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<string>", "..."],
  "weaknesses": ["<string>", "..."],
  "missing_keywords": ["<keyword or skill>", "..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "section": "<resume section>", "issue": "<what is wrong>", "suggestion": "<exactly how to fix it>"}}
  ],
  "rewritten_bullets": [
    {{"original": "<bullet copied from the resume>", "improved": "<stronger version using action verbs and metrics only if supported by the resume>"}}
  ]
}}

Give 5-8 improvements and 3-5 rewritten bullets.

{jd_block}

RESUME:
{resume_text[:MAX_RESUME_CHARS]}
"""


def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating ```json fences and stray text."""
    if not raw:
        raise ValueError("The model returned an empty response.")
    cleaned = raw.strip()
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", cleaned, flags=re.IGNORECASE)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start != -1 and end > start:
            return json.loads(cleaned[start : end + 1])
        raise ValueError("Could not parse the model response as JSON.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize_result(data: dict) -> dict:
    """Make sure every field exists with the right type so the UI never crashes."""
    cats = data.get("category_scores") or {}
    return {
        "overall_score": _clamp(data.get("overall_score")),
        "category_scores": {
            "Keywords match": _clamp(cats.get("keywords_match")),
            "Formatting & structure": _clamp(cats.get("formatting_structure")),
            "Experience impact": _clamp(cats.get("experience_impact")),
            "Skills relevance": _clamp(cats.get("skills_relevance")),
            "Readability & clarity": _clamp(cats.get("readability_clarity")),
        },
        "summary": str(data.get("summary") or ""),
        "strengths": [str(x) for x in (data.get("strengths") or [])],
        "weaknesses": [str(x) for x in (data.get("weaknesses") or [])],
        "missing_keywords": [str(x) for x in (data.get("missing_keywords") or [])],
        "improvements": [i for i in (data.get("improvements") or []) if isinstance(i, dict)],
        "rewritten_bullets": [
            b for b in (data.get("rewritten_bullets") or []) if isinstance(b, dict)
        ],
    }


def is_transient_error(exc: Exception) -> bool:
    """True for temporary Gemini errors (overloaded, rate limited, timeouts)."""
    if getattr(exc, "code", None) in (429, 500, 503, 504):
        return True
    message = str(exc).upper()
    return any(marker in message for marker in TRANSIENT_MARKERS)


def _generate_with_retry(client, model, prompt, config):
    """Call Gemini, retrying temporary failures with exponential backoff."""
    for attempt in range(RETRIES_PER_MODEL):
        try:
            return client.models.generate_content(model=model, contents=prompt, config=config)
        except Exception as exc:
            if not is_transient_error(exc) or attempt == RETRIES_PER_MODEL - 1:
                raise
            time.sleep(BACKOFF_SECONDS * (2 ** attempt))


def analyze_with_gemini(api_key: str, model: str, resume_text: str, job_description: str) -> dict:
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)
    config = types.GenerateContentConfig(
        temperature=0.2,
        response_mime_type="application/json",
    )
    prompt = build_prompt(resume_text, job_description)

    models_to_try = [model] + [m for m in FALLBACK_MODELS if m != model]
    last_error = None
    for current_model in models_to_try:
        try:
            for _ in range(2):  # one retry if the JSON comes back malformed
                response = _generate_with_retry(client, current_model, prompt, config)
                try:
                    return normalize_result(parse_json_response(response.text))
                except ValueError as exc:
                    last_error = exc
        except Exception as exc:
            last_error = exc
            if not is_transient_error(exc):
                raise  # e.g. invalid API key or unknown model: fallback won't help
    raise last_error


def get_api_key(sidebar_value: str) -> str:
    if sidebar_value.strip():
        return sidebar_value.strip()
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass
    return os.environ.get("GEMINI_API_KEY", "")


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


# --------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------
def main():
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.caption("Upload your resume and get an ATS score with practical improvements.")

    with st.sidebar:
        st.header("Settings")
        key_input = st.text_input(
            "Gemini API key",
            type="password",
            help="Leave empty if GEMINI_API_KEY is set in Streamlit secrets or environment.",
        )
        model = st.text_input("Gemini model", value=DEFAULT_MODEL)
        st.markdown("Get a free key at [Google AI Studio](https://aistudio.google.com/apikey).")
        st.info("Your resume is sent to the Gemini API for analysis. Don't upload anything you aren't comfortable sharing.")

    col_left, col_right = st.columns(2)
    with col_left:
        uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    with col_right:
        job_description = st.text_area(
            "Job description (optional, improves keyword matching)", height=170
        )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        api_key = get_api_key(key_input)
        if not api_key:
            st.error("Please enter your Gemini API key in the sidebar.")
            st.stop()

        try:
            resume_text = extract_text(uploaded.getvalue(), uploaded.name)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            st.stop()

        if len(resume_text) < 100:
            st.error(
                "Very little text could be extracted. If your resume is a scanned image, "
                "ATS systems can't read it either - export a text-based PDF or DOCX instead."
            )
            st.stop()

        with st.spinner("Analyzing your resume..."):
            try:
                ai = analyze_with_gemini(api_key, model, resume_text, job_description)
            except Exception as exc:
                if is_transient_error(exc):
                    st.error(
                        "Google's Gemini service is overloaded right now (the app already "
                        "retried and tried a backup model). Please wait a minute and click "
                        "Analyze again."
                    )
                else:
                    st.error(f"Analysis failed: {exc}")
                st.stop()
            local = local_checks(resume_text)

        # Final score: 70% AI judgement + 30% rule-based format checks
        final = round(0.7 * ai["overall_score"] + 0.3 * local["score"])
        st.session_state["report"] = {"ai": ai, "local": local, "final": final}

    report = st.session_state.get("report")
    if not report:
        st.stop()

    ai, local, final = report["ai"], report["local"], report["final"]

    st.divider()
    m1, m2, m3 = st.columns(3)
    m1.metric("Final ATS score", f"{final}/100", score_label(final))
    m2.metric("AI content score", f"{ai['overall_score']}/100")
    m3.metric("Format checks", f"{local['score']}/100")
    st.progress(final / 100)
    if ai["summary"]:
        st.write(ai["summary"])

    tab1, tab2, tab3, tab4 = st.tabs(
        ["Improvements", "Scores & checks", "Keywords", "Rewritten bullets"]
    )

    with tab1:
        if ai["improvements"]:
            order = {"high": 0, "medium": 1, "low": 2}
            items = sorted(
                ai["improvements"],
                key=lambda i: order.get(str(i.get("priority", "")).lower(), 3),
            )
            for item in items:
                priority = item.get("priority", "")
                icon = {"high": "🔴", "medium": "🟠", "low": "🟢"}.get(str(priority).lower(), "⚪")
                with st.expander(f"{icon} {priority} - {item.get('section', 'General')}"):
                    st.markdown(f"**Issue:** {item.get('issue', '')}")
                    st.markdown(f"**Fix:** {item.get('suggestion', '')}")
        else:
            st.write("No improvements returned.")
        c1, c2 = st.columns(2)
        with c1:
            st.subheader("✅ Strengths")
            for s in ai["strengths"]:
                st.markdown(f"- {s}")
        with c2:
            st.subheader("⚠️ Weaknesses")
            for w in ai["weaknesses"]:
                st.markdown(f"- {w}")

    with tab2:
        st.subheader("AI category scores")
        for name, val in ai["category_scores"].items():
            st.write(f"{name}: **{val}/100**")
            st.progress(val / 100)
        st.subheader("Format checks")
        for c in local["checks"]:
            st.markdown(f"{'✅' if c['passed'] else '❌'} {c['label']}")

    with tab3:
        if ai["missing_keywords"]:
            st.write("Consider adding these (only if they truthfully apply to you):")
            st.write(", ".join(f"`{k}`" for k in ai["missing_keywords"]))
        else:
            st.write("No missing keywords reported.")

    with tab4:
        if ai["rewritten_bullets"]:
            for b in ai["rewritten_bullets"]:
                st.markdown(f"**Before:** {b.get('original', '')}")
                st.markdown(f"**After:** {b.get('improved', '')}")
                st.divider()
        else:
            st.write("No rewritten bullets returned.")

    st.download_button(
        "Download report (JSON)",
        data=json.dumps({"final_score": final, **ai, "format_checks": local}, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


if __name__ == "__main__":
    main()
